
import os
import sys
import numpy as np
import pandas as pd
from scipy.stats import ttest_rel
from sklearn.metrics import roc_auc_score, accuracy_score, classification_report, confusion_matrix
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Seed for exact reproducibility
np.random.seed(42)

def main():
    repo_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    model_dir = os.path.join(repo_root, "Model")
    excel_root = os.path.join(repo_root, "Model_Results_Excel")
    
    # Destination directories
    desktop_dir = r"C:\Users\IHCK\Desktop\Dataset_1_Bootstrap_Results"
    dataset1_dir = os.path.join(model_dir, "Dataset_1")
    excel_out_dir = os.path.join(excel_root, "02_Bootstrap_and_Statistical_Tests", "Dataset_1_Combined")
    
    os.makedirs(desktop_dir, exist_ok=True)
    os.makedirs(dataset1_dir, exist_ok=True)
    os.makedirs(excel_out_dir, exist_ok=True)
    
    dataset_name = "Dataset_1"
    sides = ["left", "right"]
    
    # 7 model architectures
    models_config = [
        ("MLP", os.path.join(model_dir, "Dataset_1", "{side}", "MLP", "results", "test_predictions.npz")),
        ("MobileNet", os.path.join(model_dir, "Dataset_1", "{side}", "MobileNet", "results", "test_predictions.npz")),
        ("PointNet", os.path.join(model_dir, "PointNet_Results", "{side}", "results", "test_predictions.npz")),
        ("ResNet+AE", os.path.join(model_dir, "Dataset_1", "{side}", "ResNet", "results", "test_predictions_ae.npz")),
        ("ResNet", os.path.join(model_dir, "Dataset_1", "{side}", "ResNet", "results", "test_predictions.npz")),
        ("SqueezeNet", os.path.join(model_dir, "Dataset_1", "{side}", "SqueezeNet", "results", "test_predictions.npz")),
        ("SVM", os.path.join(model_dir, "Dataset_1", "{side}", "SVM", "results", "test_predictions.npz")),
    ]
    
    print("=" * 80)
    print("GENERATING COMBINED BOOTSTRAP RESULTS FOR DATASET 1 (ALL 7 MODELS INCL. POINTNET)")
    print("=" * 80)
    
    # 1. Build Model Legend
    model_entries = []
    idx = 0
    for side in sides:
        for m_name, path_pattern in models_config:
            npz_path = path_pattern.format(side=side)
            if not os.path.exists(npz_path):
                print(f"[ERROR] Prediction file not found: {npz_path}")
                sys.exit(1)
            model_tag = f"{dataset_name}_{side}_{m_name}"
            model_entries.append({
                "index": idx,
                "model_tag": model_tag,
                "dataset": dataset_name,
                "side": side,
                "model_name": m_name,
                "npz_path": npz_path
            })
            idx += 1
            
    print(f"Total evaluated models: {len(model_entries)} (7 Left, 7 Right)")
    
    # Generate Legend text
    legend_lines = ["Index | Model Name", "-" * 50]
    for me in model_entries:
        legend_lines.append(f"{me['index']} : {me['model_tag']}")
    legend_text = "\n".join(legend_lines) + "\n"
    
    # 2. Run 1,000 Paired Bootstrap Resamplings
    N_BOOT = 1000
    bootstrap_columns = {}
    bootstrap_columns["round"] = list(range(1, N_BOOT + 1))
    
    final_summary_rows = []
    perf_summary_rows = []
    
    for side in sides:
        cohort_models = [m for m in model_entries if m["side"] == side]
        first_npz = np.load(cohort_models[0]["npz_path"])
        yt_first = first_npz["y_test"] if "y_test" in first_npz else first_npz["y_true"]
        N = len(yt_first)
        
        # Shared random bootstrap indices for true paired comparison
        boot_indices = np.random.randint(0, N, size=(N_BOOT, N))
        print(f"\nProcessing Side: {side.upper()} (N_test={N}, Models={len(cohort_models)})...")
        
        for me in cohort_models:
            i = me["index"]
            data = np.load(me["npz_path"])
            y_true = data["y_test"] if "y_test" in data else data["y_true"]
            y_prob = data["y_prob"]
            if y_prob.ndim == 2:
                y_prob = y_prob[:, 1] if y_prob.shape[1] == 2 else y_prob.flatten()
            y_pred = (y_prob >= 0.5).astype(int)
            
            # Full test metrics (point estimates)
            cm_full = confusion_matrix(y_true, y_pred, labels=[0, 1])
            tn_f, fp_f, fn_f, tp_f = cm_full.ravel()
            acc_full = accuracy_score(y_true, y_pred)
            sens_c0_f = tn_f / (tn_f + fp_f) if (tn_f + fp_f) > 0 else 0.0
            spec_c0_f = tp_f / (tp_f + fn_f) if (tp_f + fn_f) > 0 else 0.0
            prec_c0_f = tn_f / (tn_f + fn_f) if (tn_f + fn_f) > 0 else 0.0
            f1_c0_f = 2 * prec_c0_f * sens_c0_f / (prec_c0_f + sens_c0_f) if (prec_c0_f + sens_c0_f) > 0 else 0.0
            
            sens_c1_f = tp_f / (tp_f + fn_f) if (tp_f + fn_f) > 0 else 0.0
            spec_c1_f = tn_f / (tn_f + fp_f) if (tn_f + fp_f) > 0 else 0.0
            prec_c1_f = tp_f / (tp_f + fp_f) if (tp_f + fp_f) > 0 else 0.0
            f1_c1_f = 2 * prec_c1_f * sens_c1_f / (prec_c1_f + sens_c1_f) if (prec_c1_f + sens_c1_f) > 0 else 0.0
            auc_full = roc_auc_score(y_true, y_prob) if len(np.unique(y_true)) > 1 else 0.5
            
            tp_arr = np.zeros(N_BOOT, dtype=int)
            tn_arr = np.zeros(N_BOOT, dtype=int)
            fp_arr = np.zeros(N_BOOT, dtype=int)
            fn_arr = np.zeros(N_BOOT, dtype=int)
            acc_arr = np.zeros(N_BOOT, dtype=float)
            sens_arr = np.zeros(N_BOOT, dtype=float)
            spec_arr = np.zeros(N_BOOT, dtype=float)
            f1_arr = np.zeros(N_BOOT, dtype=float)
            auc_arr = np.zeros(N_BOOT, dtype=float)
            
            sens_c0_arr = np.zeros(N_BOOT, dtype=float)
            spec_c0_arr = np.zeros(N_BOOT, dtype=float)
            f1_c0_arr = np.zeros(N_BOOT, dtype=float)
            
            for b_idx in range(N_BOOT):
                sample_idx = boot_indices[b_idx]
                yt = y_true[sample_idx]
                yp = y_pred[sample_idx]
                ypr = y_prob[sample_idx]
                
                tp = int(np.sum((yp == 1) & (yt == 1)))
                tn = int(np.sum((yp == 0) & (yt == 0)))
                fp = int(np.sum((yp == 1) & (yt == 0)))
                fn = int(np.sum((yp == 0) & (yt == 1)))
                
                tp_arr[b_idx] = tp
                tn_arr[b_idx] = tn
                fp_arr[b_idx] = fp
                fn_arr[b_idx] = fn
                
                acc = (tp + tn) / N
                acc_arr[b_idx] = round(acc, 4)
                
                sens = tp / (tp + fn) if (tp + fn) > 0 else 0.0
                sens_arr[b_idx] = round(sens, 4)
                
                spec = tn / (tn + fp) if (tn + fp) > 0 else 0.0
                spec_arr[b_idx] = round(spec, 4)
                
                denom1 = 2 * tp + fp + fn
                f1 = (2 * tp) / denom1 if denom1 > 0 else 0.0
                f1_arr[b_idx] = round(f1, 4)
                
                sens_c0 = spec
                spec_c0 = sens
                denom0 = 2 * tn + fp + fn
                f1_c0 = (2 * tn) / denom0 if denom0 > 0 else 0.0
                sens_c0_arr[b_idx] = sens_c0
                spec_c0_arr[b_idx] = spec_c0
                f1_c0_arr[b_idx] = f1_c0
                
                if len(np.unique(yt)) > 1:
                    try:
                        auc_val = roc_auc_score(yt, ypr)
                    except ValueError:
                        auc_val = 0.5
                else:
                    auc_val = 0.5
                auc_arr[b_idx] = round(auc_val, 4)
                
            # Populate wide table matching example schema exactly
            bootstrap_columns[f"TP_{i}"] = tp_arr
            bootstrap_columns[f"TN_{i}"] = tn_arr
            bootstrap_columns[f"FP_{i}"] = fp_arr
            bootstrap_columns[f"FN_{i}"] = fn_arr
            bootstrap_columns[f"Accuracy_{i}"] = acc_arr
            bootstrap_columns[f"Sensitivity_{i}"] = sens_arr
            bootstrap_columns[f"Specificity_{i}"] = spec_arr
            bootstrap_columns[f"F1_{i}"] = f1_arr
            bootstrap_columns[f"AUC_{i}"] = auc_arr
            
            # Summary stats
            final_summary_rows.append({
                "Dataset": me["dataset"],
                "Side": me["side"],
                "Model": me["model_name"],
                "Mean_Accuracy": round(float(np.mean(acc_arr)), 4),
                "SD_Accuracy": round(float(np.std(acc_arr)), 4),
                "Mean_Sensitivity_Class0": round(float(np.mean(sens_c0_arr)), 4),
                "SD_Sensitivity_Class0": round(float(np.std(sens_c0_arr)), 4),
                "Mean_Specificity_Class0": round(float(np.mean(spec_c0_arr)), 4),
                "SD_Specificity_Class0": round(float(np.std(spec_c0_arr)), 4),
                "Mean_F1_Class0": round(float(np.mean(f1_c0_arr)), 4),
                "SD_F1_Class0": round(float(np.std(f1_c0_arr)), 4),
                "Mean_Sensitivity_Class1": round(float(np.mean(sens_arr)), 4),
                "SD_Sensitivity_Class1": round(float(np.std(sens_arr)), 4),
                "Mean_Specificity_Class1": round(float(np.mean(spec_arr)), 4),
                "SD_Specificity_Class1": round(float(np.std(spec_arr)), 4),
                "Mean_F1_Class1": round(float(np.mean(f1_arr)), 4),
                "SD_F1_Class1": round(float(np.std(f1_arr)), 4),
                "Mean_AUC": round(float(np.mean(auc_arr)), 4),
                "SD_AUC": round(float(np.std(auc_arr)), 4),
            })
            
            perf_summary_rows.append({
                "Dataset": me["dataset"],
                "Side": me["side"].capitalize(),
                "Model": me["model_name"],
                "Overall Accuracy": round(acc_full, 4),
                "Overall Mean ± SD": f"{np.mean(acc_arr):.4f} ± {np.std(acc_arr):.4f}",
                "Class 0 Acc": round(sens_c0_f, 4),
                "Class 0 Precision": round(prec_c0_f, 4),
                "Class 0 Recall": round(sens_c0_f, 4),
                "Class 0 F1": round(f1_c0_f, 4),
                "Class 0 Mean ± SD": f"{np.mean(f1_c0_arr):.4f} ± {np.std(f1_c0_arr):.4f}",
                "Class 1 Acc": round(sens_c1_f, 4),
                "Class 1 Precision": round(prec_c1_f, 4),
                "Class 1 Recall": round(sens_c1_f, 4),
                "Class 1 F1": round(f1_c1_f, 4),
                "Class 1 Mean ± SD": f"{np.mean(f1_arr):.4f} ± {np.std(f1_arr):.4f}",
                "ROC AUC": round(auc_full, 4)
            })
            
    df_combined_boot = pd.DataFrame(bootstrap_columns)
    print(f"\n[OK] Combined Bootstrap Table shape: {df_combined_boot.shape} (1000 rounds, {df_combined_boot.shape[1]} columns)")
    
    df_final_summary = pd.DataFrame(final_summary_rows)
    df_perf_summary = pd.DataFrame(perf_summary_rows)
    
    # 3. Compute Pairwise P-Values (Paired t-tests)
    metrics_to_test = [
        ("Accuracy", "Accuracy"),
        ("AUC", "AUC"),
        ("F1_Class0", "F1_C0"),
        ("F1_Class1", "F1"),
        ("Sensitivity_Class0", "Specificity"),
        ("Sensitivity_Class1", "Sensitivity"),
        ("Specificity_Class0", "Sensitivity"),
        ("Specificity_Class1", "Specificity")
    ]
    
    pvalue_rows = []
    for metric_name, col_prefix in metrics_to_test:
        for side in sides:
            cohort_models = [m for m in model_entries if m["side"] == side]
            n_m = len(cohort_models)
            for a_idx in range(n_m):
                for b_idx in range(a_idx + 1, n_m):
                    mA = cohort_models[a_idx]
                    mB = cohort_models[b_idx]
                    idxA = mA["index"]
                    idxB = mB["index"]
                    
                    if col_prefix == "F1_C0":
                        tnA = df_combined_boot[f"TN_{idxA}"]
                        fpA = df_combined_boot[f"FP_{idxA}"]
                        fnA = df_combined_boot[f"FN_{idxA}"]
                        denomA = 2 * tnA + fpA + fnA
                        seriesA = np.where(denomA > 0, 2 * tnA / denomA, 0.0)
                        
                        tnB = df_combined_boot[f"TN_{idxB}"]
                        fpB = df_combined_boot[f"FP_{idxB}"]
                        fnB = df_combined_boot[f"FN_{idxB}"]
                        denomB = 2 * tnB + fpB + fnB
                        seriesB = np.where(denomB > 0, 2 * tnB / denomB, 0.0)
                    else:
                        seriesA = df_combined_boot[f"{col_prefix}_{idxA}"].values
                        seriesB = df_combined_boot[f"{col_prefix}_{idxB}"].values
                        
                    meanA = float(np.mean(seriesA))
                    meanB = float(np.mean(seriesB))
                    
                    if np.all(seriesA == seriesB):
                        p_val = 1.0
                    else:
                        t_stat, p_val = ttest_rel(seriesA, seriesB, nan_policy="omit")
                        if np.isnan(p_val):
                            p_val = 1.0
                            
                    if p_val < 0.01:
                        sig = "Yes (p<0.01)"
                    elif p_val < 0.05:
                        sig = "Yes (p<0.05)"
                    else:
                        sig = "No"
                        
                    pvalue_rows.append({
                        "Metric": metric_name,
                        "Dataset": dataset_name,
                        "Side": side.capitalize(),
                        "Model_A": mA["model_name"],
                        "Model_B": mB["model_name"],
                        "Mean_Value_A": round(meanA, 4),
                        "Mean_Value_B": round(meanB, 4),
                        "P-Value": p_val,
                        "Significant": sig
                    })
                    
    df_pvalue = pd.DataFrame(pvalue_rows)
    df_non_sig = df_pvalue[df_pvalue["Significant"] == "No"].copy()
    print(f"[OK] Computed {len(df_pvalue)} pairwise tests ({len(df_non_sig)} non-significant)")
    
    # 4. Save CSVs to all target directories
    target_dirs = [desktop_dir, dataset1_dir, excel_out_dir]
    
    for t_dir in target_dirs:
        # Combined Bootstrap CSV
        c_path = os.path.join(t_dir, "Combined_Bootstrap_Results.csv")
        df_combined_boot.to_csv(c_path, index=False)
        
        # Legend TXT
        l_path = os.path.join(t_dir, "Model_Legend.txt")
        with open(l_path, "w", encoding="utf-8") as f:
            f.write(legend_text)
            
        # Final Summary Mean SD CSV
        s_path = os.path.join(t_dir, "Final_Summary_Mean_SD.csv")
        df_final_summary.to_csv(s_path, index=False)
        
        # P-Value Summary CSV
        p_path = os.path.join(t_dir, "PValue_Significance_Summary.csv")
        df_pvalue.to_csv(p_path, index=False)
        
        # Non-Significant CSV
        ns_path = os.path.join(t_dir, "Non_Significant_Pairs_Summary.csv")
        df_non_sig.to_csv(ns_path, index=False)
        
        # Performance Summary CSV
        perf_path = os.path.join(t_dir, "Model_Performance_Summary.csv")
        df_perf_summary.to_csv(perf_path, index=False)
        
        print(f"Saved complete suite to: {t_dir}")
        
    # 5. Build Master Styled Excel Workbook
    wb = openpyxl.Workbook()
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=13, bold=True, color="FFFFFF")
    section_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    header_font = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    bold_font = Font(name=font_family, size=9, bold=True)
    normal_font = Font(name=font_family, size=9)
    
    title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")       # Dark Navy
    side_banner_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid") # Medium Navy
    info_fill = PatternFill(start_color="41719C", end_color="41719C", fill_type="solid")        # Slate
    overall_fill = PatternFill(start_color="333F48", end_color="333F48", fill_type="solid")     # Charcoal
    c0_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")          # Steel Blue
    c1_fill = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")          # Warm Rust
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    best_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")        # Mint Green
    
    thin_border = Border(
        left=Side(style="thin", color="D0D7DE"), right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"), bottom=Side(style="thin", color="D0D7DE")
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    # TAB 1: Model_Performance_Summary
    ws1 = wb.active
    ws1.title = "Performance_Summary"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.merge_cells("A1:P1")
    t1 = ws1["A1"]
    t1.value = "Dataset 1 Model Performance Summary (Class 0: Healthy vs Class 1: Epilepsy TLE - 7 Models)"
    t1.font = title_font
    t1.fill = title_fill
    t1.alignment = align_center
    ws1.row_dimensions[1].height = 34
    
    # Headers
    ws1.merge_cells("A3:C3")
    ws1["A3"] = "Model Information"
    ws1["A3"].font = header_font
    ws1["A3"].fill = info_fill
    ws1["A3"].alignment = align_center
    
    ws1.merge_cells("D3:E3")
    ws1["D3"] = "Overall Metrics"
    ws1["D3"].font = header_font
    ws1["D3"].fill = overall_fill
    ws1["D3"].alignment = align_center
    
    ws1.merge_cells("F3:J3")
    ws1["F3"] = "Class 0: Healthy Control (HC)"
    ws1["F3"].font = header_font
    ws1["F3"].fill = c0_fill
    ws1["F3"].alignment = align_center
    
    ws1.merge_cells("K3:O3")
    ws1["K3"] = "Class 1: Epilepsy / TLE"
    ws1["K3"].font = header_font
    ws1["K3"].fill = c1_fill
    ws1["K3"].alignment = align_center
    
    ws1.cell(row=3, column=16, value="Discrimination").fill = info_fill
    ws1.cell(row=3, column=16).font = header_font
    ws1.cell(row=3, column=16).alignment = align_center
    
    sub_headers = [
        "Dataset", "Side", "Model",
        "Overall Accuracy", "Overall Mean ± SD",
        "Class 0 Acc", "Class 0 Precision", "Class 0 Recall", "Class 0 F1", "Class 0 Mean ± SD",
        "Class 1 Acc", "Class 1 Precision", "Class 1 Recall", "Class 1 F1", "Class 1 Mean ± SD",
        "ROC AUC"
    ]
    ws1.row_dimensions[4].height = 24
    for ci, h in enumerate(sub_headers, 1):
        c = ws1.cell(row=4, column=ci, value=h)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    for r_i in [3, 4]:
        for c_i in range(1, 17):
            ws1.cell(row=r_i, column=c_i).border = thin_border
            
    curr_row = 5
    for side in sides:
        sub_df = df_perf_summary[df_perf_summary["Side"] == side.capitalize()]
        best_acc = sub_df["Overall Accuracy"].max()
        
        ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=16)
        sb = ws1.cell(row=curr_row, column=1, value=f"Dataset: Dataset_1 | Side: {side.upper()} HIPPOCAMPUS")
        sb.font = section_font
        sb.fill = side_banner_fill
        sb.alignment = align_left
        ws1.row_dimensions[curr_row].height = 22
        curr_row += 1
        
        for _, row in sub_df.iterrows():
            ws1.row_dimensions[curr_row].height = 20
            is_best = (row["Overall Accuracy"] == best_acc)
            fill_to_use = best_fill if is_best else (zebra_fill if curr_row % 2 == 0 else PatternFill(fill_type=None))
            
            vals = [
                row["Dataset"], row["Side"], row["Model"],
                f"{row['Overall Accuracy']*100:.2f}%", row["Overall Mean ± SD"],
                f"{row['Class 0 Acc']*100:.2f}%", f"{row['Class 0 Precision']*100:.2f}%", f"{row['Class 0 Recall']*100:.2f}%", f"{row['Class 0 F1']:.4f}", row["Class 0 Mean ± SD"],
                f"{row['Class 1 Acc']*100:.2f}%", f"{row['Class 1 Precision']*100:.2f}%", f"{row['Class 1 Recall']*100:.2f}%", f"{row['Class 1 F1']:.4f}", row["Class 1 Mean ± SD"],
                f"{row['ROC AUC']:.4f}"
            ]
            for col_idx, val in enumerate(vals, 1):
                cell = ws1.cell(row=curr_row, column=col_idx, value=val)
                cell.font = bold_font if (is_best and col_idx == 4) else normal_font
                if fill_to_use.fill_type is not None:
                    cell.fill = fill_to_use
                cell.border = thin_border
                cell.alignment = align_center if col_idx in [1, 2, 4, 6, 7, 8, 9, 11, 12, 13, 14, 16] else (align_left if col_idx == 3 else align_center)
            curr_row += 1
            
    # TAB 2: Model_Legend
    ws2 = wb.create_sheet(title="Model_Legend")
    ws2.views.sheetView[0].showGridLines = True
    ws2.merge_cells("A1:C1")
    ws2["A1"] = "Model Index Legend (Mapping for Combined Bootstrap Table)"
    ws2["A1"].font = title_font
    ws2["A1"].fill = title_fill
    ws2["A1"].alignment = align_center
    ws2.row_dimensions[1].height = 30
    
    headers_l = ["Index", "Model Name / Tag", "Side"]
    ws2.row_dimensions[3].height = 22
    for ci, h in enumerate(headers_l, 1):
        c = ws2.cell(row=3, column=ci, value=h)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    for r_i, me in enumerate(model_entries, 4):
        ws2.row_dimensions[r_i].height = 19
        c1 = ws2.cell(row=r_i, column=1, value=me["index"])
        c2 = ws2.cell(row=r_i, column=2, value=me["model_tag"])
        c3 = ws2.cell(row=r_i, column=3, value=me["side"].capitalize())
        for c, al in zip([c1, c2, c3], [align_center, align_left, align_center]):
            c.font = normal_font
            c.border = thin_border
            c.alignment = al
            if r_i % 2 == 0:
                c.fill = zebra_fill
                
    # TAB 3: Final_Summary_Mean_SD
    ws3 = wb.create_sheet(title="Final_Summary_Mean_SD")
    ws3.views.sheetView[0].showGridLines = True
    ws3.merge_cells("A1:S1")
    ws3["A1"] = "Final Bootstrap Summary: Mean and Standard Deviation (1,000 Rounds)"
    ws3["A1"].font = title_font
    ws3["A1"].fill = title_fill
    ws3["A1"].alignment = align_center
    ws3.row_dimensions[1].height = 30
    
    s_cols = list(df_final_summary.columns)
    ws3.row_dimensions[3].height = 24
    for ci, h in enumerate(s_cols, 1):
        c = ws3.cell(row=3, column=ci, value=h)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    for r_i, r_data in enumerate(df_final_summary.itertuples(index=False), 4):
        ws3.row_dimensions[r_i].height = 19
        for c_i, val in enumerate(r_data, 1):
            c = ws3.cell(row=r_i, column=c_i, value=val)
            c.font = normal_font
            c.border = thin_border
            c.alignment = align_center if c_i != 3 else align_left
            if r_i % 2 == 0:
                c.fill = zebra_fill

    # TAB 4: PValue_Summary
    ws4 = wb.create_sheet(title="PValue_Summary")
    ws4.views.sheetView[0].showGridLines = True
    ws4.merge_cells("A1:I1")
    ws4["A1"] = "Pairwise Statistical Significance (Paired t-tests across 1,000 Bootstrap Rounds)"
    ws4["A1"].font = title_font
    ws4["A1"].fill = title_fill
    ws4["A1"].alignment = align_center
    ws4.row_dimensions[1].height = 30
    
    p_cols = list(df_pvalue.columns)
    ws4.row_dimensions[3].height = 24
    for ci, h in enumerate(p_cols, 1):
        c = ws4.cell(row=3, column=ci, value=h)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    for r_i, r_data in enumerate(df_pvalue.itertuples(index=False), 4):
        ws4.row_dimensions[r_i].height = 18
        for c_i, val in enumerate(r_data, 1):
            c = ws4.cell(row=r_i, column=c_i, value=val)
            c.font = normal_font
            c.border = thin_border
            c.alignment = align_center
            if c_i == 9: # Significant
                if "p<0.01" in str(val):
                    c.fill = best_fill
                    c.font = bold_font
                elif "p<0.05" in str(val):
                    c.fill = zebra_fill
            elif r_i % 2 == 0:
                c.fill = zebra_fill
                
    # TAB 5: Non_Significant_Pairs
    ws5 = wb.create_sheet(title="Non_Significant_Pairs")
    ws5.views.sheetView[0].showGridLines = True
    ws5.merge_cells("A1:I1")
    ws5["A1"] = "Pairs with No Statistically Significant Difference (p >= 0.05)"
    ws5["A1"].font = title_font
    ws5["A1"].fill = overall_fill
    ws5["A1"].alignment = align_center
    ws5.row_dimensions[1].height = 30
    
    ws5.row_dimensions[3].height = 24
    for ci, h in enumerate(p_cols, 1):
        c = ws5.cell(row=3, column=ci, value=h)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    for r_i, r_data in enumerate(df_non_sig.itertuples(index=False), 4):
        ws5.row_dimensions[r_i].height = 18
        for c_i, val in enumerate(r_data, 1):
            c = ws5.cell(row=r_i, column=c_i, value=val)
            c.font = normal_font
            c.border = thin_border
            c.alignment = align_center
            if r_i % 2 == 0:
                c.fill = zebra_fill
                
    # Auto-adjust column widths across all sheets
    for ws in [ws1, ws2, ws3, ws4, ws5]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2]:
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)
            
    # Save Workbook to target destinations
    excel_filename = "Dataset_1_Model_Performance_and_Bootstrap_Master.xlsx"
    for t_dir in target_dirs:
        wb_path = os.path.join(t_dir, excel_filename)
        wb.save(wb_path)
        print(f"Saved Excel Workbook: {wb_path}")
        
    print("\n" + "=" * 80)
    print("ALL COMBINED BOOTSTRAP FILES SUCCESSFULLY GENERATED!")
    print("=" * 80)

if __name__ == "__main__":
    main()
