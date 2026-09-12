import os
import numpy as np
import pandas as pd
from sklearn.metrics import classification_report, accuracy_score, roc_auc_score
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Set seed for exact reproducibility
np.random.seed(42)

def compute_metrics_per_class(y_true, y_prob, n_boot=1000):
    if y_prob.ndim == 2:
        y_prob = y_prob[:, 1] if y_prob.shape[1] == 2 else y_prob.flatten()
    y_pred = (y_prob > 0.5).astype(int)
    n = len(y_true)
    
    # Sklearn classification report
    cr = classification_report(y_true, y_pred, output_dict=True, zero_division=0)
    c0 = cr.get('0', {'precision': 0, 'recall': 0, 'f1-score': 0, 'support': int(np.sum(y_true == 0))})
    c1 = cr.get('1', {'precision': 0, 'recall': 0, 'f1-score': 0, 'support': int(np.sum(y_true == 1))})
    
    # Overall metrics
    overall_acc = accuracy_score(y_true, y_pred) * 100
    test_auc = roc_auc_score(y_true, y_prob) if len(np.unique(y_true)) > 1 else 0.5
    
    # Bootstrap 1,000 iterations for Mean +- SD of Overall, Class 0, and Class 1
    boot_indices = np.random.randint(0, n, size=(n_boot, n))
    yt_b = y_true[boot_indices]
    yp_b = y_pred[boot_indices]
    
    # Overall Accuracy Bootstrap
    boot_accs = np.mean(yt_b == yp_b, axis=1) * 100
    mean_acc = float(np.mean(boot_accs))
    sd_acc = float(np.std(boot_accs))
    
    # Class 0 (Healthy) Accuracy / Recall Bootstrap
    tp0 = np.sum((yp_b == 0) & (yt_b == 0), axis=1)
    fn0 = np.sum((yp_b == 1) & (yt_b == 0), axis=1)
    tot0 = tp0 + fn0
    recs0 = np.divide(tp0, tot0, out=np.zeros(n_boot), where=(tot0 > 0)) * 100
    c0_mean_sd = f"{float(np.mean(recs0)):.2f}% ± {float(np.std(recs0)):.2f}%"
    
    # Class 1 (Epilepsy) Accuracy / Recall Bootstrap
    tp1 = np.sum((yp_b == 1) & (yt_b == 1), axis=1)
    fn1 = np.sum((yp_b == 0) & (yt_b == 1), axis=1)
    tot1 = tp1 + fn1
    recs1 = np.divide(tp1, tot1, out=np.zeros(n_boot), where=(tot1 > 0)) * 100
    c1_mean_sd = f"{float(np.mean(recs1)):.2f}% ± {float(np.std(recs1)):.2f}%"
    
    return {
        "Overall_Acc": overall_acc,
        "Overall_Mean_SD": f"{mean_acc:.2f}% ± {sd_acc:.2f}%",
        "ROC_AUC": test_auc,
        
        # Class 0: Healthy Control
        "C0_Acc": c0['recall'] * 100,      # Accuracy on Healthy samples (% correct)
        "C0_Prec": c0['precision'] * 100,  # Precision when predicting Healthy
        "C0_Rec": c0['recall'] * 100,      # Recall / Specificity
        "C0_F1": c0['f1-score'] * 100,
        "C0_Mean_SD": c0_mean_sd,
        "C0_Support": int(c0['support']),
        
        # Class 1: Epilepsy / TLE Patient
        "C1_Acc": c1['recall'] * 100,      # Accuracy on Epilepsy samples (% correct)
        "C1_Prec": c1['precision'] * 100,  # Precision when predicting Epilepsy
        "C1_Rec": c1['recall'] * 100,      # Recall / Sensitivity
        "C1_F1": c1['f1-score'] * 100,
        "C1_Mean_SD": c1_mean_sd,
        "C1_Support": int(c1['support']),
    }

def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    excel_root = os.path.abspath(os.path.join(script_dir, ".."))
    repo_root = os.path.abspath(os.path.join(excel_root, ".."))
    model_dir = os.path.join(repo_root, "Model")
    
    dir_workbooks = os.path.join(excel_root, "01_Excel_Workbooks")
    dir_summaries = os.path.join(excel_root, "03_Performance_Summary_CSVs")
    os.makedirs(dir_workbooks, exist_ok=True)
    os.makedirs(dir_summaries, exist_ok=True)
    
    datasets = ["Ds005602", "All_Augment_tain", "Ds004469"]
    sides = ["left", "right"]
    
    models_config = [
        ("SVM", "SVM", "test_predictions.npz"),
        ("MLP", "MLP", "test_predictions.npz"),
        ("ResNet (Standard)", "ResNet", "test_predictions.npz"),
        ("ResNet (AutoEncoder)", "ResNet", "test_predictions_ae.npz"),
        ("MobileNet", "MobileNet", "test_predictions.npz"),
        ("SqueezeNet", "SqueezeNet", "test_predictions.npz"),
        ("PointNet", "PointNet", "test_predictions.npz")
    ]
    
    rows = []
    
    for ds in datasets:
        for side in sides:
            for m_name, subfolder, npz_name in models_config:
                npz_path = os.path.join(model_dir, ds, side, subfolder, "results", npz_name)
                if not os.path.exists(npz_path):
                    print(f"[WARN] Not found: {npz_path}")
                    continue
                
                npz = np.load(npz_path)
                y_true = npz['y_test'] if 'y_test' in npz else npz['y_true']
                y_prob = npz['y_prob']
                
                m = compute_metrics_per_class(y_true, y_prob, n_boot=1000)
                
                rows.append({
                    "Dataset": ds,
                    "Side": side.capitalize(),
                    "Model": m_name,
                    
                    # Overall
                    "Overall_Accuracy": f"{m['Overall_Acc']:.2f}%",
                    "Overall_Mean_SD": m['Overall_Mean_SD'],
                    
                    # Class 0: Healthy Control (has its own acc, precision, recall, mean+sd)
                    "Class0_Healthy_Acc": f"{m['C0_Acc']:.2f}%",
                    "Class0_Healthy_Precision": f"{m['C0_Prec']:.2f}%",
                    "Class0_Healthy_Recall": f"{m['C0_Rec']:.2f}%",
                    "Class0_Healthy_F1": f"{m['C0_F1']:.2f}%",
                    "Class0_Healthy_Mean_SD": m['C0_Mean_SD'],
                    
                    # Class 1: Epilepsy TLE (has its own acc, precision, recall, mean+sd)
                    "Class1_Epilepsy_Acc": f"{m['C1_Acc']:.2f}%",
                    "Class1_Epilepsy_Precision": f"{m['C1_Prec']:.2f}%",
                    "Class1_Epilepsy_Recall": f"{m['C1_Rec']:.2f}%",
                    "Class1_Epilepsy_F1": f"{m['C1_F1']:.2f}%",
                    "Class1_Epilepsy_Mean_SD": m['C1_Mean_SD'],
                    
                    "ROC_AUC": f"{m['ROC_AUC']:.4f}",
                    "_raw_acc": m['Overall_Acc']
                })
                
    df = pd.DataFrame([{k: v for k, v in r.items() if not k.startswith('_')} for r in rows])
    
    # Save to the single CSV file
    csv_path = os.path.join(dir_summaries, "model_performance_summary.csv")
    try:
        df.to_csv(csv_path, index=False, encoding="utf-8-sig")
        print(f"[OK] Saved single CSV: {csv_path}")
    except PermissionError:
        alt_csv = os.path.join(dir_summaries, "model_performance_summary_updated.csv")
        df.to_csv(alt_csv, index=False, encoding="utf-8-sig")
        print(f"[NOTE] model_performance_summary.csv is locked in Excel. Saved updated CSV to: {alt_csv}")

    # ==========================================================
    # Build Single Master Excel Workbook
    # ==========================================================
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model_Performance_Summary"
    ws.views.sheetView[0].showGridLines = True
    
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=13, bold=True, color="FFFFFF")
    section_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    header_font = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    bold_font = Font(name=font_family, size=9, bold=True)
    normal_font = Font(name=font_family, size=9)
    
    title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")       # Dark Navy
    side_banner_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid") # Medium Navy
    info_fill = PatternFill(start_color="41719C", end_color="41719C", fill_type="solid")        # Slate
    overall_fill = PatternFill(start_color="333F48", end_color="333F48", fill_type="solid")     # Charcoal
    c0_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")          # Steel Blue (Healthy)
    c1_fill = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")          # Warm Rust (Epilepsy)
    
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    best_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")        # Mint Green
    
    thin_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE")
    )
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    # Title
    ws.merge_cells("A1:P1")
    t = ws["A1"]
    t.value = "Model Performance Summary: Evaluated by Class (Class 0: Healthy vs Class 1: Epilepsy TLE)"
    t.font = title_font
    t.fill = title_fill
    t.alignment = align_center
    ws.row_dimensions[1].height = 34
    
    # Two-tier header:
    # Row 3: Group categories
    ws.row_dimensions[3].height = 20
    ws.merge_cells("A3:C3")
    ws["A3"] = "Model Information"
    ws["A3"].font = header_font
    ws["A3"].fill = info_fill
    ws["A3"].alignment = align_center
    
    ws.merge_cells("D3:E3")
    ws["D3"] = "Overall Metrics"
    ws["D3"].font = header_font
    ws["D3"].fill = overall_fill
    ws["D3"].alignment = align_center
    
    ws.merge_cells("F3:J3")
    ws["F3"] = "Class 0: Healthy Control (คนปกติ)"
    ws["F3"].font = header_font
    ws["F3"].fill = c0_fill
    ws["F3"].alignment = align_center
    
    ws.merge_cells("K3:O3")
    ws["K3"] = "Class 1: Epilepsy / TLE (ผู้ป่วย)"
    ws["K3"].font = header_font
    ws["K3"].fill = c1_fill
    ws["K3"].alignment = align_center
    
    ws.cell(row=3, column=16, value="Discrimination").fill = info_fill
    ws.cell(row=3, column=16).font = header_font
    ws.cell(row=3, column=16).alignment = align_center
    
    # Row 4: Sub-columns
    sub_headers = [
        "Dataset", "Side", "Model (Row)",
        "Accuracy (acc)", "Mean ± SD",
        "Accuracy (acc)", "Precision (prec)", "Recall (spec)", "F1-Score", "Mean ± SD",
        "Accuracy (acc)", "Precision (prec)", "Recall (sens)", "F1-Score", "Mean ± SD",
        "ROC AUC"
    ]
    ws.row_dimensions[4].height = 24
    for ci, h in enumerate(sub_headers, 1):
        cell = ws.cell(row=4, column=ci, value=h)
        cell.font = header_font
        cell.fill = info_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    for r_i in [3, 4]:
        for c_i in range(1, 17):
            ws.cell(row=r_i, column=c_i).border = thin_border
            
    curr_row = 5
    for ds in datasets:
        for side in sides:
            subset = [r for r in rows if r['Dataset'] == ds and r['Side'] == side.capitalize()]
            if not subset:
                continue
                
            # Section Banner
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=16)
            sb = ws.cell(row=curr_row, column=1, value=f"📍 Dataset: {ds} | Side: {side.upper()} HIPPOCAMPUS")
            sb.font = section_font
            sb.fill = side_banner_fill
            sb.alignment = align_left
            ws.row_dimensions[curr_row].height = 22
            curr_row += 1
            
            best_acc = max(r['_raw_acc'] for r in subset)
            
            for idx, r in enumerate(subset):
                ws.row_dimensions[curr_row].height = 19
                is_best = (r['_raw_acc'] == best_acc)
                bg = best_fill if is_best else (zebra_fill if idx % 2 == 1 else None)
                
                vals = [
                    r['Dataset'], r['Side'], r['Model'],
                    r['Overall_Accuracy'], r['Overall_Mean_SD'],
                    r['Class0_Healthy_Acc'], r['Class0_Healthy_Precision'], r['Class0_Healthy_Recall'], r['Class0_Healthy_F1'], r['Class0_Healthy_Mean_SD'],
                    r['Class1_Epilepsy_Acc'], r['Class1_Epilepsy_Precision'], r['Class1_Epilepsy_Recall'], r['Class1_Epilepsy_F1'], r['Class1_Epilepsy_Mean_SD'],
                    r['ROC_AUC']
                ]
                for ci, val in enumerate(vals, 1):
                    cell = ws.cell(row=curr_row, column=ci, value=val)
                    cell.font = bold_font if is_best else normal_font
                    if bg: cell.fill = bg
                    cell.border = thin_border
                    cell.alignment = align_left if ci in [1, 3] else align_center
                curr_row += 1
            curr_row += 1 # small spacing between dataset/side
            
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to single Excel file
    xlsx_path = os.path.join(dir_workbooks, "Model_Performance_Summary.xlsx")
    wb.save(xlsx_path)
    print(f"[OK] Saved single Excel file: {xlsx_path}")

if __name__ == "__main__":
    main()
