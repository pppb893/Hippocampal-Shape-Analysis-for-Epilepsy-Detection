import os
import sys
import numpy as np
import pandas as pd
from scipy.stats import ttest_rel
from sklearn.metrics import roc_auc_score, accuracy_score, classification_report
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure exact reproducibility
np.random.seed(42)

def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    excel_root = os.path.abspath(os.path.join(script_dir, ".."))
    repo_root = os.path.abspath(os.path.join(excel_root, ".."))
    model_dir = os.path.join(repo_root, "Model")
    
    dir_workbooks = os.path.join(excel_root, "01_Excel_Workbooks")
    dir_bootstrap = os.path.join(excel_root, "02_Bootstrap_and_Statistical_Tests")
    dir_summaries = os.path.join(excel_root, "03_Performance_Summary_CSVs")
    
    os.makedirs(dir_workbooks, exist_ok=True)
    os.makedirs(dir_bootstrap, exist_ok=True)
    os.makedirs(dir_summaries, exist_ok=True)
    
    # 3 cohorts, 2 sides
    datasets = ["All_Augment_tain", "Ds005602", "Ds004469"]
    sides = ["left", "right"]
    
    # 7 model architectures
    models_config = [
        ("MLP", "MLP", "test_predictions.npz"),
        ("MobileNet", "MobileNet", "test_predictions.npz"),
        ("PointNet", "PointNet", "test_predictions.npz"),
        ("ResNet+AE", "ResNet", "test_predictions_ae.npz"),
        ("ResNet", "ResNet", "test_predictions.npz"),
        ("SqueezeNet", "SqueezeNet", "test_predictions.npz"),
        ("SVM", "SVM", "test_predictions.npz"),
    ]
    
    print("=" * 70)
    print("GENERATING NEWEST BOOTSTRAP EVALUATION RESULTS")
    print("Includes Ds004469, SqueezeNet, PointNet across all 42 models")
    print("=" * 70)
    
    # 1. Build Model Legend
    model_entries = []
    idx = 0
    for ds in datasets:
        for side in sides:
            for m_name, subfolder, npz_name in models_config:
                model_tag = f"{ds}_{side}_{m_name}"
                npz_path = os.path.join(model_dir, ds, side, subfolder, "results", npz_name)
                if not os.path.exists(npz_path):
                    print(f"[ERROR] Prediction file missing: {npz_path}")
                    sys.exit(1)
                model_entries.append({
                    "index": idx,
                    "model_tag": model_tag,
                    "dataset": ds,
                    "side": side,
                    "model_name": m_name,
                    "npz_path": npz_path
                })
                idx += 1
                
    total_models = len(model_entries)
    print(f"Total evaluated models: {total_models}")
    
    # Save Model Legend TXT files
    legend_lines = ["Index | Model Name", "-" * 50]
    for me in model_entries:
        legend_lines.append(f"{me['index']} : {me['model_tag']}")
    legend_text = "\n".join(legend_lines) + "\n"
    
    legend_path = os.path.join(dir_bootstrap, "Model_Legend.txt")
    with open(legend_path, "w", encoding="utf-8") as f:
        f.write(legend_text)
    print(f"[OK] Saved Legend: {legend_path}")

    # 2. Run 1,000 Paired Bootstrap Resamplings
    N_BOOT = 1000
    bootstrap_columns = {}
    bootstrap_columns["round"] = list(range(1, N_BOOT + 1))
    
    # Storage for aggregated stats per model
    final_summary_rows = []
    
    # Group by cohort and side for paired sampling
    for ds in datasets:
        for side in sides:
            cohort_models = [m for m in model_entries if m["dataset"] == ds and m["side"] == side]
            if not cohort_models:
                continue
                
            # Load first model to get N and test set labels
            first_npz = np.load(cohort_models[0]["npz_path"])
            yt_first = first_npz["y_test"] if "y_test" in first_npz else first_npz["y_true"]
            N = len(yt_first)
            
            # Generate shared bootstrap sample indices for this cohort & side
            # Shape: (N_BOOT, N)
            boot_indices = np.random.randint(0, N, size=(N_BOOT, N))
            
            print(f"Processing Cohort: {ds} | Side: {side.upper()} (N_test={N}, Models={len(cohort_models)})...")
            
            for me in cohort_models:
                i = me["index"]
                data = np.load(me["npz_path"])
                y_true = data["y_test"] if "y_test" in data else data["y_true"]
                y_prob = data["y_prob"]
                if y_prob.ndim == 2:
                    y_prob = y_prob[:, 1] if y_prob.shape[1] == 2 else y_prob.flatten()
                y_pred = (y_prob >= 0.5).astype(int)
                
                # Arrays to hold 1,000 round results
                tp_arr = np.zeros(N_BOOT, dtype=int)
                tn_arr = np.zeros(N_BOOT, dtype=int)
                fp_arr = np.zeros(N_BOOT, dtype=int)
                fn_arr = np.zeros(N_BOOT, dtype=int)
                acc_arr = np.zeros(N_BOOT, dtype=float)
                sens_arr = np.zeros(N_BOOT, dtype=float)
                spec_arr = np.zeros(N_BOOT, dtype=float)
                f1_arr = np.zeros(N_BOOT, dtype=float)
                auc_arr = np.zeros(N_BOOT, dtype=float)
                
                # Class 0 metrics per round for final summary
                sens_c0_arr = np.zeros(N_BOOT, dtype=float)
                spec_c0_arr = np.zeros(N_BOOT, dtype=float)
                f1_c0_arr = np.zeros(N_BOOT, dtype=float)
                
                for b_idx in range(N_BOOT):
                    sample_idx = boot_indices[b_idx]
                    yt = y_true[sample_idx]
                    yp = y_pred[sample_idx]
                    ypr = y_prob[sample_idx]
                    
                    tp = int(np.sum((yp == 1) & (yt == 1)))
                    tn = int(np.sum((yp == 0) & (yt == 0)))
                    fp = int(np.sum((yp == 1) & (yt == 0)))
                    fn = int(np.sum((yp == 0) & (yt == 1)))
                    
                    tp_arr[b_idx] = tp
                    tn_arr[b_idx] = tn
                    fp_arr[b_idx] = fp
                    fn_arr[b_idx] = fn
                    
                    # Accuracy
                    acc = (tp + tn) / N
                    acc_arr[b_idx] = round(acc, 4)
                    
                    # Sensitivity (Class 1 Recall / Sensitivity)
                    sens = tp / (tp + fn) if (tp + fn) > 0 else 0.0
                    sens_arr[b_idx] = round(sens, 4)
                    
                    # Specificity (Class 0 Recall / Specificity)
                    spec = tn / (tn + fp) if (tn + fp) > 0 else 0.0
                    spec_arr[b_idx] = round(spec, 4)
                    
                    # F1 Score (Class 1)
                    denom1 = 2 * tp + fp + fn
                    f1 = (2 * tp) / denom1 if denom1 > 0 else 0.0
                    f1_arr[b_idx] = round(f1, 4)
                    
                    # Class 0 metrics (Healthy Control)
                    sens_c0 = spec
                    spec_c0 = sens
                    denom0 = 2 * tn + fp + fn
                    f1_c0 = (2 * tn) / denom0 if denom0 > 0 else 0.0
                    sens_c0_arr[b_idx] = sens_c0
                    spec_c0_arr[b_idx] = spec_c0
                    f1_c0_arr[b_idx] = f1_c0
                    
                    # AUC
                    if len(np.unique(yt)) > 1:
                        try:
                            auc_val = roc_auc_score(yt, ypr)
                        except ValueError:
                            auc_val = 0.5
                    else:
                        auc_val = 0.5
                    auc_arr[b_idx] = round(auc_val, 4)
                    
                # Save columns matching example schema
                bootstrap_columns[f"TP_{i}"] = tp_arr
                bootstrap_columns[f"TN_{i}"] = tn_arr
                bootstrap_columns[f"FP_{i}"] = fp_arr
                bootstrap_columns[f"FN_{i}"] = fn_arr
                bootstrap_columns[f"Accuracy_{i}"] = acc_arr
                bootstrap_columns[f"Sensitivity_{i}"] = sens_arr
                bootstrap_columns[f"Specificity_{i}"] = spec_arr
                bootstrap_columns[f"F1_{i}"] = f1_arr
                bootstrap_columns[f"AUC_{i}"] = auc_arr
                
                # Summary stats for this model
                final_summary_rows.append({
                    "Dataset": me["dataset"],
                    "Side": me["side"],
                    "Model": me["model_name"],
                    "Mean_Accuracy": round(float(np.mean(acc_arr)), 4),
                    "SD_Accuracy": round(float(np.std(acc_arr)), 4),
                    "Mean_Sensitivity_Class0": round(float(np.mean(sens_c0_arr)), 4),
                    "SD_Sensitivity_Class0": round(float(np.std(sens_c0_arr)), 4),
                    "Mean_Specificity_Class0": round(float(np.mean(spec_c0_arr)), 4),
                    "SD_Specificity_Class0": round(float(np.std(spec_c0_arr)), 4),
                    "Mean_F1_Class0": round(float(np.mean(f1_c0_arr)), 4),
                    "SD_F1_Class0": round(float(np.std(f1_c0_arr)), 4),
                    "Mean_Sensitivity_Class1": round(float(np.mean(sens_arr)), 4),
                    "SD_Sensitivity_Class1": round(float(np.std(sens_arr)), 4),
                    "Mean_Specificity_Class1": round(float(np.mean(spec_arr)), 4),
                    "SD_Specificity_Class1": round(float(np.std(spec_arr)), 4),
                    "Mean_F1_Class1": round(float(np.mean(f1_arr)), 4),
                    "SD_F1_Class1": round(float(np.std(f1_arr)), 4),
                    "Mean_AUC": round(float(np.mean(auc_arr)), 4),
                    "SD_AUC": round(float(np.std(auc_arr)), 4),
                })
                
    # 3. Create Combined Bootstrap DataFrame and Save CSVs
    df_combined_boot = pd.DataFrame(bootstrap_columns)
    print(f"\nCombined Bootstrap DataFrame Shape: {df_combined_boot.shape} (Expected: 1000 rows, 379 cols)")
    
    boot_csv = os.path.join(dir_bootstrap, "Combined_Bootstrap_Results.csv")
    df_combined_boot.to_csv(boot_csv, index=False)
    print(f"[OK] Saved Combined Bootstrap CSV: {boot_csv}")

    # 4. Create Final Summary Mean SD DataFrame and Save CSVs
    df_final_summary = pd.DataFrame(final_summary_rows)
    summary_csv = os.path.join(dir_bootstrap, "Final_Summary_Mean_SD.csv")
    df_final_summary.to_csv(summary_csv, index=False)
    print(f"[OK] Saved Final Summary Mean SD CSV: {summary_csv}")

    # 5. Compute P-Value Summary (Pairwise Paired t-tests)
    # Metrics to test
    metrics_to_test = [
        ("Accuracy", "Accuracy"),
        ("AUC", "AUC"),
        ("F1_Class0", "F1_C0"),
        ("F1_Class1", "F1"),
        ("Sensitivity_Class0", "Specificity"), # Recall for C0 is Specificity
        ("Sensitivity_Class1", "Sensitivity"), # Recall for C1 is Sensitivity
        ("Specificity_Class0", "Sensitivity"),
        ("Specificity_Class1", "Specificity")
    ]
    
    pvalue_rows = []
    
    for metric_name, col_prefix in metrics_to_test:
        for ds in datasets:
            for side in sides:
                cohort_models = [m for m in model_entries if m["dataset"] == ds and m["side"] == side]
                n_m = len(cohort_models)
                for a_idx in range(n_m):
                    for b_idx in range(a_idx + 1, n_m):
                        mA = cohort_models[a_idx]
                        mB = cohort_models[b_idx]
                        idxA = mA["index"]
                        idxB = mB["index"]
                        
                        if col_prefix == "F1_C0":
                            # Compute F1 for Class 0 directly from TN, FP, FN
                            tnA = df_combined_boot[f"TN_{idxA}"]
                            fpA = df_combined_boot[f"FP_{idxA}"]
                            fnA = df_combined_boot[f"FN_{idxA}"]
                            seriesA = np.where((2 * tnA + fpA + fnA) > 0, 2 * tnA / (2 * tnA + fpA + fnA), 0.0)
                            
                            tnB = df_combined_boot[f"TN_{idxB}"]
                            fpB = df_combined_boot[f"FP_{idxB}"]
                            fnB = df_combined_boot[f"FN_{idxB}"]
                            seriesB = np.where((2 * tnB + fpB + fnB) > 0, 2 * tnB / (2 * tnB + fpB + fnB), 0.0)
                        else:
                            seriesA = df_combined_boot[f"{col_prefix}_{idxA}"].values
                            seriesB = df_combined_boot[f"{col_prefix}_{idxB}"].values
                            
                        meanA = float(np.mean(seriesA))
                        meanB = float(np.mean(seriesB))
                        
                        # Paired t-test
                        if np.all(seriesA == seriesB):
                            p_val = 1.0
                        else:
                            t_stat, p_val = ttest_rel(seriesA, seriesB, nan_policy="omit")
                            if np.isnan(p_val):
                                p_val = 1.0
                                
                        if p_val < 0.01:
                            sig = "Yes (p<0.01)"
                        elif p_val < 0.05:
                            sig = "Yes (p<0.05)"
                        else:
                            sig = "No"
                            
                        pvalue_rows.append({
                            "Metric": metric_name,
                            "Dataset": ds,
                            "Side": side,
                            "Model_A": mA["model_name"],
                            "Model_B": mB["model_name"],
                            "Mean_Value_A": round(meanA, 4),
                            "Mean_Value_B": round(meanB, 4),
                            "P-Value": p_val,
                            "Significant": sig
                        })
                        
    df_pvalue = pd.DataFrame(pvalue_rows)
    pvalue_csv = os.path.join(dir_bootstrap, "PValue_Significance_Summary.csv")
    df_pvalue.to_csv(pvalue_csv, index=False)
    print(f"[OK] Saved P-Value Summary CSV ({len(df_pvalue)} comparisons): {pvalue_csv}")

    # 6. Build Master Multi-Tab Excel Spreadsheet
    # Combining Model_Performance_Summary, Final_Summary_Mean_SD, PValue_Summary, Model_Legend, and Combined_Bootstrap
    wb = openpyxl.Workbook()
    
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=13, bold=True, color="FFFFFF")
    section_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    header_font = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    bold_font = Font(name=font_family, size=9, bold=True)
    normal_font = Font(name=font_family, size=9)
    
    title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")       # Dark Navy
    side_banner_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid") # Medium Navy
    info_fill = PatternFill(start_color="41719C", end_color="41719C", fill_type="solid")        # Slate
    overall_fill = PatternFill(start_color="333F48", end_color="333F48", fill_type="solid")     # Charcoal
    c0_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")          # Steel Blue (Healthy)
    c1_fill = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")          # Warm Rust (Epilepsy)
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    best_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")        # Mint Green
    
    thin_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE")
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    # -----------------------------------------------------------
    # SHEET 1: Model_Performance_Summary
    # -----------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Model_Performance_Summary"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.merge_cells("A1:P1")
    t1 = ws1["A1"]
    t1.value = "Model Performance Summary: Evaluated by Class (Class 0: Healthy vs Class 1: Epilepsy TLE)"
    t1.font = title_font
    t1.fill = title_fill
    t1.alignment = align_center
    ws1.row_dimensions[1].height = 34
    
    # Tier 1 Header
    ws1.row_dimensions[3].height = 20
    ws1.merge_cells("A3:C3")
    ws1["A3"] = "Model Information"
    ws1["A3"].font = header_font
    ws1["A3"].fill = info_fill
    ws1["A3"].alignment = align_center
    
    ws1.merge_cells("D3:E3")
    ws1["D3"] = "Overall Metrics"
    ws1["D3"].font = header_font
    ws1["D3"].fill = overall_fill
    ws1["D3"].alignment = align_center
    
    ws1.merge_cells("F3:J3")
    ws1["F3"] = "Class 0: Healthy Control (HC)"
    ws1["F3"].font = header_font
    ws1["F3"].fill = c0_fill
    ws1["F3"].alignment = align_center
    
    ws1.merge_cells("K3:O3")
    ws1["K3"] = "Class 1: Epilepsy / TLE"
    ws1["K3"].font = header_font
    ws1["K3"].fill = c1_fill
    ws1["K3"].alignment = align_center
    
    ws1.cell(row=3, column=16, value="Discrimination").fill = info_fill
    ws1.cell(row=3, column=16).font = header_font
    ws1.cell(row=3, column=16).alignment = align_center
    
    # Tier 2 Sub-headers
    sub_headers = [
        "Dataset", "Side", "Model",
        "Overall Accuracy", "Overall Mean ± SD",
        "Class 0 Acc", "Class 0 Precision", "Class 0 Recall", "Class 0 F1", "Class 0 Mean ± SD",
        "Class 1 Acc", "Class 1 Precision", "Class 1 Recall", "Class 1 F1", "Class 1 Mean ± SD",
        "ROC AUC"
    ]
    ws1.row_dimensions[4].height = 24
    for ci, h in enumerate(sub_headers, 1):
        c = ws1.cell(row=4, column=ci, value=h)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    for r_i in [3, 4]:
        for c_i in range(1, 17):
            ws1.cell(row=r_i, column=c_i).border = thin_border

    # Populate rows
    curr_row = 5
    summary_rows_for_csv = []
    
    for ds in datasets:
        for side in sides:
            cohort_models = [m for m in model_entries if m["dataset"] == ds and m["side"] == side]
            if not cohort_models:
                continue
                
            # Section Banner
            ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=16)
            sb = ws1.cell(row=curr_row, column=1, value=f"Dataset: {ds} | Side: {side.upper()} HIPPOCAMPUS")
            sb.font = section_font
            sb.fill = side_banner_fill
            sb.alignment = align_left
            ws1.row_dimensions[curr_row].height = 22
            curr_row += 1
            
            # Find best raw accuracy
            raw_accs = []
            for me in cohort_models:
                d = np.load(me["npz_path"])
                yt = d["y_test"] if "y_test" in d else d["y_true"]
                yp = d["y_prob"]
                if yp.ndim == 2:
                    yp = yp[:, 1] if yp.shape[1] == 2 else yp.flatten()
                pred = (yp >= 0.5).astype(int)
                raw_accs.append(accuracy_score(yt, pred))
            best_acc = max(raw_accs)
            
            for idx_m, me in enumerate(cohort_models):
                ws1.row_dimensions[curr_row].height = 19
                i = me["index"]
                d = np.load(me["npz_path"])
                yt = d["y_test"] if "y_test" in d else d["y_true"]
                yp = d["y_prob"]
                if yp.ndim == 2:
                    yp = yp[:, 1] if yp.shape[1] == 2 else yp.flatten()
                pred = (yp >= 0.5).astype(int)
                
                cr = classification_report(yt, pred, output_dict=True, zero_division=0)
                c0 = cr.get('0', {'precision': 0, 'recall': 0, 'f1-score': 0})
                c1 = cr.get('1', {'precision': 0, 'recall': 0, 'f1-score': 0})
                
                raw_acc = accuracy_score(yt, pred) * 100
                auc_score = roc_auc_score(yt, yp) if len(np.unique(yt)) > 1 else 0.5
                
                acc_series = df_combined_boot[f"Accuracy_{i}"] * 100
                spec_series = df_combined_boot[f"Specificity_{i}"] * 100
                sens_series = df_combined_boot[f"Sensitivity_{i}"] * 100
                
                overall_mean_sd = f"{np.mean(acc_series):.2f}% ± {np.std(acc_series):.2f}%"
                c0_mean_sd = f"{np.mean(spec_series):.2f}% ± {np.std(spec_series):.2f}%"
                c1_mean_sd = f"{np.mean(sens_series):.2f}% ± {np.std(sens_series):.2f}%"
                
                is_best = (accuracy_score(yt, pred) == best_acc)
                bg = best_fill if is_best else (zebra_fill if idx_m % 2 == 1 else None)
                
                row_vals = [
                    me["dataset"],
                    me["side"].capitalize(),
                    me["model_name"],
                    f"{raw_acc:.2f}%",
                    overall_mean_sd,
                    f"{c0['recall']*100:.2f}%",
                    f"{c0['precision']*100:.2f}%",
                    f"{c0['recall']*100:.2f}%",
                    f"{c0['f1-score']*100:.2f}%",
                    c0_mean_sd,
                    f"{c1['recall']*100:.2f}%",
                    f"{c1['precision']*100:.2f}%",
                    f"{c1['recall']*100:.2f}%",
                    f"{c1['f1-score']*100:.2f}%",
                    c1_mean_sd,
                    f"{auc_score:.4f}"
                ]
                
                summary_rows_for_csv.append({
                    "Dataset": me["dataset"],
                    "Side": me["side"].capitalize(),
                    "Model": me["model_name"],
                    "Overall_Accuracy": f"{raw_acc:.2f}%",
                    "Overall_Mean_SD": overall_mean_sd,
                    "Class0_Healthy_Acc": f"{c0['recall']*100:.2f}%",
                    "Class0_Healthy_Precision": f"{c0['precision']*100:.2f}%",
                    "Class0_Healthy_Recall": f"{c0['recall']*100:.2f}%",
                    "Class0_Healthy_F1": f"{c0['f1-score']*100:.2f}%",
                    "Class0_Healthy_Mean_SD": c0_mean_sd,
                    "Class1_Epilepsy_Acc": f"{c1['recall']*100:.2f}%",
                    "Class1_Epilepsy_Precision": f"{c1['precision']*100:.2f}%",
                    "Class1_Epilepsy_Recall": f"{c1['recall']*100:.2f}%",
                    "Class1_Epilepsy_F1": f"{c1['f1-score']*100:.2f}%",
                    "Class1_Epilepsy_Mean_SD": c1_mean_sd,
                    "ROC_AUC": f"{auc_score:.4f}"
                })
                
                for ci, val in enumerate(row_vals, 1):
                    c = ws1.cell(row=curr_row, column=ci, value=val)
                    c.font = bold_font if is_best else normal_font
                    if bg: c.fill = bg
                    c.border = thin_border
                    c.alignment = align_left if ci in [1, 3] else align_center
                curr_row += 1
            curr_row += 1 # separation space
            
    for col in ws1.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 13)

    # -----------------------------------------------------------
    # SHEET 2: Final_Summary_Mean_SD
    # -----------------------------------------------------------
    ws2 = wb.create_sheet(title="Final_Summary_Mean_SD")
    ws2.views.sheetView[0].showGridLines = True
    
    headers_s2 = list(df_final_summary.columns)
    ws2.row_dimensions[1].height = 24
    for ci, h in enumerate(headers_s2, 1):
        c = ws2.cell(row=1, column=ci, value=h)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    for ri, row in df_final_summary.iterrows():
        ws2.row_dimensions[ri + 2].height = 19
        bg = zebra_fill if ri % 2 == 1 else None
        for ci, h in enumerate(headers_s2, 1):
            val = row[h]
            c = ws2.cell(row=ri + 2, column=ci, value=val)
            c.font = normal_font
            if bg: c.fill = bg
            c.border = thin_border
            c.alignment = align_left if ci in [1, 3] else align_center
            
    for col in ws2.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # -----------------------------------------------------------
    # SHEET 3: PValue_Summary
    # -----------------------------------------------------------
    ws3 = wb.create_sheet(title="PValue_Summary")
    ws3.views.sheetView[0].showGridLines = True
    
    headers_s3 = list(df_pvalue.columns)
    ws3.row_dimensions[1].height = 24
    for ci, h in enumerate(headers_s3, 1):
        c = ws3.cell(row=1, column=ci, value=h)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    sig_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # subtle light yellow
    for ri, row in df_pvalue.iterrows():
        ws3.row_dimensions[ri + 2].height = 18
        is_sig = "Yes" in str(row["Significant"])
        bg = sig_fill if is_sig else (zebra_fill if ri % 2 == 1 else None)
        for ci, h in enumerate(headers_s3, 1):
            val = row[h]
            c = ws3.cell(row=ri + 2, column=ci, value=val)
            c.font = bold_font if is_sig and h in ["P-Value", "Significant"] else normal_font
            if bg: c.fill = bg
            c.border = thin_border
            c.alignment = align_left if ci in [1, 2, 4, 5] else align_center
            
    for col in ws3.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 13)

    # -----------------------------------------------------------
    # SHEET 4: Model_Legend
    # -----------------------------------------------------------
    ws4 = wb.create_sheet(title="Model_Legend")
    ws4.views.sheetView[0].showGridLines = True
    
    ws4.row_dimensions[1].height = 24
    ws4.cell(row=1, column=1, value="Index").font = header_font
    ws4.cell(row=1, column=1).fill = info_fill
    ws4.cell(row=1, column=1).alignment = align_center
    ws4.cell(row=1, column=1).border = thin_border
    
    ws4.cell(row=1, column=2, value="Model Name / Tag").font = header_font
    ws4.cell(row=1, column=2).fill = info_fill
    ws4.cell(row=1, column=2).alignment = align_center
    ws4.cell(row=1, column=2).border = thin_border
    
    for ri, me in enumerate(model_entries):
        ws4.row_dimensions[ri + 2].height = 19
        bg = zebra_fill if ri % 2 == 1 else None
        c1 = ws4.cell(row=ri + 2, column=1, value=me["index"])
        c2 = ws4.cell(row=ri + 2, column=2, value=me["model_tag"])
        c1.font = normal_font
        c2.font = normal_font
        if bg:
            c1.fill = bg
            c2.fill = bg
        c1.border = thin_border
        c2.border = thin_border
        c1.alignment = align_center
        c2.alignment = align_left
        
    ws4.column_dimensions["A"].width = 12
    ws4.column_dimensions["B"].width = 40

    # -----------------------------------------------------------
    # SHEET 5: Combined_Bootstrap_Results
    # -----------------------------------------------------------
    ws5 = wb.create_sheet(title="Combined_Bootstrap_Results")
    ws5.views.sheetView[0].showGridLines = True
    
    boot_cols = list(df_combined_boot.columns)
    ws5.row_dimensions[1].height = 24
    for ci, col_name in enumerate(boot_cols, 1):
        c = ws5.cell(row=1, column=ci, value=col_name)
        c.font = header_font
        c.fill = info_fill
        c.alignment = align_center
        c.border = thin_border
        
    print("Writing Combined Bootstrap Results to Excel Sheet 5...")
    boot_vals = df_combined_boot.values
    for ri in range(len(boot_vals)):
        if (ri + 1) % 250 == 0:
            print(f"  ... written {ri + 1}/{len(boot_vals)} bootstrap rounds")
        for ci in range(len(boot_cols)):
            ws5.cell(row=ri + 2, column=ci + 1, value=boot_vals[ri, ci])
            
    # Save Excel Workbooks
    master_xlsx_path = os.path.join(dir_workbooks, "Model_Performance_and_Bootstrap_Master.xlsx")
    wb.save(master_xlsx_path)
    print(f"\n[OK] Saved Master Multi-Tab Excel Workbook: {master_xlsx_path}")
    
    # Also update model_performance_summary.csv in 03_Performance_Summary_CSVs
    df_perf_summary = pd.DataFrame(summary_rows_for_csv)
    perf_csv_path = os.path.join(dir_summaries, "model_performance_summary.csv")
    try:
        df_perf_summary.to_csv(perf_csv_path, index=False)
        print(f"[OK] Updated model_performance_summary.csv: {perf_csv_path}")
    except PermissionError:
        alt_csv = os.path.join(dir_summaries, "model_performance_summary_updated.csv")
        df_perf_summary.to_csv(alt_csv, index=False)
        print(f"[NOTE] Saved to alt CSV: {alt_csv}")
        
    # Also generate the standalone clean Model_Performance_Summary.xlsx in 01_Excel_Workbooks
    try:
        from . import generate_model_excel
        generate_model_excel.main()
    except Exception:
        try:
            import generate_model_excel
            generate_model_excel.main()
        except Exception as e:
            print(f"[WARN] Error updating standalone Model_Performance_Summary.xlsx: {e}")
        
    print("\n" + "=" * 70)
    print("ALL FILES GENERATED SUCCESSFULLY!")
    print("=" * 70)

if __name__ == "__main__":
    main()
